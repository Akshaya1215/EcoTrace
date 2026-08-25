from fastapi import APIRouter, HTTPException, Query, UploadFile, File, Form
from datetime import datetime
from bson import ObjectId
from typing import List, Optional
from app.database import db, clean_db_doc, clean_db_list

router = APIRouter(prefix="/listings", tags=["Listings"])

# Status values:
# DRAFT, AI_ANALYZED, AVAILABLE, ACCEPTED, PICKUP_SCHEDULED, COLLECTED, PROCESSING, MATERIAL_RECOVERED, PAYMENT_PENDING, COMPLETED, CANCELLED

@router.post("")
async def create_listing(listing: dict):
    try:
        # Assign unique id if not present
        if "_id" not in listing:
            listing["_id"] = str(ObjectId())
            
        listing["status"] = listing.get("status", "AVAILABLE")
        listing["createdAt"] = listing.get("createdAt", datetime.utcnow().isoformat())
        listing["recyclerId"] = None
        listing["recoveredMaterials"] = None
        listing["transactionId"] = None
        listing["pickupDetails"] = None
        
        # Save listing
        res = db.listings.insert_one(listing)
        return {"id": res.inserted_id, "status": listing["status"], "success": True}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to create listing: {str(e)}")

@router.get("")
async def get_listings(status: Optional[str] = None):
    try:
        query = {}
        if status:
            query["status"] = status
        
        listings = db.listings.find(query)
        # Sort by createdAt descending
        listings = sorted(listings, key=lambda x: x.get("createdAt", ""), reverse=True)
        return clean_db_list(listings)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to fetch listings: {str(e)}")

@router.get("/seller/{seller_id}")
async def get_seller_listings(seller_id: str):
    try:
        listings = db.listings.find({"sellerId": seller_id})
        listings = sorted(listings, key=lambda x: x.get("createdAt", ""), reverse=True)
        return clean_db_list(listings)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to fetch seller listings: {str(e)}")

@router.get("/available")
async def get_available_listings():
    try:
        # Available means status is AVAILABLE
        listings = db.listings.find({"status": "AVAILABLE"})
        listings = sorted(listings, key=lambda x: x.get("createdAt", ""), reverse=True)
        return clean_db_list(listings)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to fetch available listings: {str(e)}")

@router.get("/accepted/{recycler_id}")
async def get_accepted_listings(recycler_id: str):
    try:
        # Get listings accepted by a recycler (accepted, pickup_scheduled, collected, processing, etc.)
        listings = db.listings.find({
            "recyclerId": recycler_id,
            "status": {"$in": ["ACCEPTED", "PICKUP_SCHEDULED", "COLLECTED", "PROCESSING", "MATERIAL_RECOVERED", "PAYMENT_PENDING", "COMPLETED"]}
        })
        listings = sorted(listings, key=lambda x: x.get("createdAt", ""), reverse=True)
        return clean_db_list(listings)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to fetch accepted listings: {str(e)}")

@router.put("/{listing_id}/accept")
async def accept_listing(listing_id: str, payload: dict):
    recycler_id = payload.get("recyclerId")
    if not recycler_id:
        raise HTTPException(status_code=400, detail="recyclerId is required")
        
    try:
        # Check if listing is AVAILABLE
        listing = db.listings.find_one({"_id": listing_id})
        if not listing:
            raise HTTPException(status_code=404, detail="Listing not found")
            
        if listing.get("status") != "AVAILABLE":
            raise HTTPException(status_code=400, detail=f"Listing is not available, current status: {listing.get('status')}")
            
        # Update listing status and recyclerId
        db.listings.update_one(
            {"_id": listing_id},
            {"$set": {
                "status": "ACCEPTED",
                "recyclerId": recycler_id,
                "acceptedAt": datetime.utcnow().isoformat()
            }}
        )
        return {"success": True, "status": "ACCEPTED"}
    except HTTPException as he:
        raise he
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to accept listing: {str(e)}")

@router.put("/{listing_id}/status")
async def update_listing_status(listing_id: str, payload: dict):
    status = payload.get("status")
    if not status:
        raise HTTPException(status_code=400, detail="status is required")
        
    try:
        # Get current listing
        listing = db.listings.find_one({"_id": listing_id})
        if not listing:
            raise HTTPException(status_code=404, detail="Listing not found")
            
        update_data = {"status": status}
        
        # Track transitions if necessary
        if status == "COMPLETED":
            update_data["completedAt"] = datetime.utcnow().isoformat()
            
        db.listings.update_one(
            {"_id": listing_id},
            {"$set": update_data}
        )
        return {"success": True, "status": status}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to update listing status: {str(e)}")

@router.post("/{listing_id}/recover-materials")
async def record_recovered_materials(listing_id: str, payload: dict):
    # payload is list of material recovery records
    recovered = payload.get("materials")
    if not recovered:
        raise HTTPException(status_code=400, detail="materials details are required")
        
    try:
        listing = db.listings.find_one({"_id": listing_id})
        if not listing:
            raise HTTPException(status_code=404, detail="Listing not found")
            
        # Calculate final recovery value based on recovered weight
        total_recovery_value = 0.0
        from app.services.valuation_service import MARKET_RATES
        
        updated_recovered_details = []
        for item in recovered:
            name = item["material"]
            rec_weight = float(item.get("recoveredWeight", 0))
            est_weight = float(item.get("estimatedWeight", 0))
            
            rate = MARKET_RATES.get(name, MARKET_RATES["Other"])
            val = rec_weight * rate
            total_recovery_value += val
            
            # calculate rate
            rec_rate = (rec_weight / est_weight) * 100 if est_weight > 0 else 100
            
            updated_recovered_details.append({
                "material": name,
                "estimatedWeight": est_weight,
                "recoveredWeight": rec_weight,
                "recoveryRate": round(rec_rate, 1),
                "estimatedValue": round(val, 2)
            })

        # Update listing status to MATERIAL_RECOVERED, store details
        db.listings.update_one(
            {"_id": listing_id},
            {"$set": {
                "status": "MATERIAL_RECOVERED",
                "recoveredMaterials": updated_recovered_details,
                "finalRecoveryValue": round(total_recovery_value, 2)
            }}
        )
        
        # Save recovered materials log
        db.materials_recovered.insert_one({
            "listingId": listing_id,
            "sellerId": listing.get("sellerId"),
            "recyclerId": listing.get("recyclerId"),
            "materials": updated_recovered_details,
            "finalRecoveryValue": round(total_recovery_value, 2),
            "createdAt": datetime.utcnow().isoformat()
        })
        
        return {
            "success": True, 
            "status": "MATERIAL_RECOVERED", 
            "finalRecoveryValue": round(total_recovery_value, 2),
            "materials": updated_recovered_details
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to record recovered materials: {str(e)}")

@router.put("/{listing_id}/cancel")
async def cancel_listing(listing_id: str):
    try:
        listing = db.listings.find_one({"_id": listing_id})
        if not listing:
            raise HTTPException(status_code=404, detail="Listing not found")
            
        if listing.get("status") not in ["AVAILABLE", "DRAFT", "AI_ANALYZED"]:
            raise HTTPException(status_code=400, detail=f"Cannot cancel listing in status {listing.get('status')}")
            
        db.listings.update_one(
            {"_id": listing_id},
            {"$set": {"status": "CANCELLED"}}
        )
        return {"success": True, "status": "CANCELLED"}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to cancel listing: {str(e)}")

@router.post("/bulk-upload")
async def bulk_upload_listings(
    file: UploadFile = File(...),
    corporateId: str = Form(...)
):
    filename = file.filename
    if not (filename.endswith(".csv") or filename.endswith(".xlsx")):
        raise HTTPException(status_code=400, detail="Only CSV and Excel (.xlsx) files are supported")
        
    try:
        import csv
        from io import StringIO, BytesIO
        from app.services.ai_service import analyze_device_data
        
        contents = await file.read()
        rows = []
        
        if filename.endswith(".csv"):
            decoded = contents.decode("utf-8-sig")
            csv_file = StringIO(decoded)
            reader = csv.DictReader(csv_file)
            for row in reader:
                rows.append(row)
        else:
            import openpyxl
            wb = openpyxl.load_workbook(BytesIO(contents))
            sheet = wb.active
            headers = [cell.value for cell in sheet[1]]
            for row_idx in range(2, sheet.max_row + 1):
                row_data = {}
                row_vals = [sheet.cell(row=row_idx, column=col_idx).value for col_idx in range(1, len(headers) + 1)]
                # Skip completely empty rows
                if not any(row_vals):
                    continue
                for idx, h in enumerate(headers):
                    if h:
                        row_data[h] = row_vals[idx]
                rows.append(row_data)
                
        total_rows = len(rows)
        valid_rows = 0
        invalid_rows_list = []
        created_listings = []
        
        # Helper to map standard names case-insensitively
        def get_value(row_dict, keys):
            for k in row_dict.keys():
                if k.lower().replace(" ", "").replace("_", "") in keys:
                    return str(row_dict[k]).strip() if row_dict[k] is not None else ""
            return ""

        seen_records = set()
        
        for idx, r in enumerate(rows):
            # Column mapping
            model_name = get_value(r, ["modelname", "devicemodel", "model", "device"])
            condition = get_value(r, ["condition", "status"]) or "good"
            location = get_value(r, ["location", "place", "address"])
            contact = get_value(r, ["contact", "contactdetails", "phone", "email"])
            description = get_value(r, ["description", "notes", "ewastedescription"])
            
            # Validation
            row_errors = []
            if not model_name:
                row_errors.append("Missing device model/name")
            if not location:
                row_errors.append("Missing location")
            if not contact:
                row_errors.append("Missing contact details")
                
            # Duplicate checking
            record_key = f"{model_name}|{location}|{contact}"
            if record_key in seen_records:
                row_errors.append("Duplicate row in upload")
            seen_records.add(record_key)
            
            if row_errors:
                invalid_rows_list.append({
                    "row": idx + 2, # 1-based, plus header row
                    "data": r,
                    "errors": row_errors
                })
                continue
                
            # Valid row - perform analysis and create listing
            try:
                analysis = analyze_device_data(model_name, description, condition)
                listing_doc = {
                    "_id": str(ObjectId()),
                    "sellerId": corporateId,
                    "deviceName": model_name,
                    "brand": analysis["device"]["brand"],
                    "category": analysis["device"]["category"],
                    "condition": condition,
                    "location": location,
                    "contactDetails": contact,
                    "aiAnalysis": analysis,
                    "estimatedWeight": analysis["estimatedWeightGrams"],
                    "materials": analysis["materials"],
                    "estimatedRecoveryValue": analysis["valuation"],
                    "environmentalImpact": analysis["environmentalImpact"],
                    "status": "AVAILABLE",
                    "createdAt": datetime.utcnow().isoformat(),
                    "recyclerId": None,
                    "recoveredMaterials": None,
                    "transactionId": None,
                    "pickupDetails": None
                }
                
                db.listings.insert_one(listing_doc)
                valid_rows += 1
                created_listings.append(listing_doc["_id"])
            except Exception as ex:
                invalid_rows_list.append({
                    "row": idx + 2,
                    "data": r,
                    "errors": [f"Error processing AI analysis: {str(ex)}"]
                })
                
        return {
            "totalRows": total_rows,
            "validRows": valid_rows,
            "invalidRowsCount": len(invalid_rows_list),
            "invalidRows": invalid_rows_list,
            "createdListingsCount": len(created_listings),
            "createdListings": created_listings,
            "success": True
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to process bulk upload: {str(e)}")

